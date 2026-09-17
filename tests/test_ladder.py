"""B3: the evidence ladder. The model proposes a level, the code sets it."""
from __future__ import annotations

import asyncio
import json

import pytest

from scan import config, io_xlsx, ladder, pipeline, sources, spec

GATES = {"adopt": {"min": 4, "or_replicated_africa_min": 3}, "adapt": {"min": 3}, "watch": {"min": 1}}


def _ev(**kw):
    base = {"method": "rct", "model_level": "E4", "independent": True, "evaluator": "J-PAL",
            "outcome_type": "outcomes", "program_matches": True, "countries": ["Kenya"],
            "method_quote": "Applicants were randomly assigned to treatment and control groups."}
    return {**base, **kw}


# --- method words ---------------------------------------------------------------------
@pytest.mark.parametrize("quote,level", [
    ("This systematic review covers 113 evaluations.", 5),
    ("Une méta-analyse de 40 études.", 5),
    ("Applicants were randomly assigned.", 4),
    ("Nous avons mené un essai randomisé.", 4),
    ("We use a difference-in-differences design.", 3),
    ("A regression discontinuity around the age cutoff.", 3),
    ("propensity score matching builds the comparison.", 3),
    ("La double différence compare les deux groupes.", 3),
    ("We compare baseline and endline surveys.", 2),
    ("A tracer study followed graduates.", 2),
    ("The program offers a matching grant to firms.", 1),
    ("The program randomly offered training to as many people as there were slots.", 4),
    ("Of these, a random half were awarded a voucher, while the other half served as the control group.", 4),
    ("We match firms claiming the subsidy with similar firms not claiming the subsidy.", 3),
    ("We use a conditional difference -in-differences approach.", 3),
    ("The analysis included pre- and post-training employment comparisons.", 2),
    ("The program trained 5,000 young people.", 1),
])
def test_method_words(quote, level):
    assert ladder.level_from_terms(quote) == level


# --- the caps -------------------------------------------------------------------------
def test_grounded_independent_rct_is_e4():
    g = ladder.grade(_ev(), implementer="Harambee", method_grounded=True, outcome_grounded=True)
    assert g["level"] == 4 and g["caps"] == [] and g["flags"] == []


def test_ungrounded_method_sentence_caps_at_e2():
    g = ladder.grade(_ev(), implementer="Harambee", method_grounded=False)
    assert g["level"] == 2 and any("not found" in c for c in g["caps"])
    assert ladder.grade(_ev(), implementer="Harambee", method_grounded=None)["level"] == 2


def test_lower_level_wins_when_model_and_code_disagree():
    g = ladder.grade(_ev(model_level="E4", method_quote="We compare baseline and endline surveys."),
                     implementer="Harambee", method_grounded=True)
    assert g["level"] == 2 and g["flags"] == ["model said E4, code set E2"]
    g2 = ladder.grade(_ev(model_level="E2"), implementer="Harambee", method_grounded=True)
    assert g2["level"] == 2


def test_method_enum_is_a_ceiling_too():
    g = ladder.grade(_ev(method="quasi_experimental"), implementer="Harambee", method_grounded=True)
    assert g["level"] == 3


def test_self_evaluation_caps_at_e3():
    g = ladder.grade(_ev(evaluator="Harambee Youth Employment Accelerator"),
                     implementer="Harambee Youth Employment Accelerator", method_grounded=True)
    assert g["level"] == 3 and any("independent" in c for c in g["caps"])
    g2 = ladder.grade(_ev(independent=False), implementer="Harambee", method_grounded=True)
    assert g2["level"] == 3


def test_outputs_only_caps_at_e2():
    g = ladder.grade(_ev(outcome_type="outputs_only"), implementer="Harambee", method_grounded=True)
    assert g["level"] == 2


def test_ungrounded_outcome_is_flagged_not_capped():
    g = ladder.grade(_ev(), implementer="Harambee", method_grounded=True, outcome_grounded=False)
    assert g["level"] == 4
    assert "outcome sentence not found word for word in the evaluation" in g["flags"]


def test_different_program_counts_for_nothing():
    g = ladder.grade(_ev(program_matches=False), implementer="Harambee", method_grounded=True)
    assert g["level"] == 1
    assert ladder.combine([g])["level"] == 1


# --- combining and replication --------------------------------------------------------
def test_two_rcts_in_different_countries_make_e5():
    a = {**_ev(countries=["Kenya"]), "level": 4}
    b = {**_ev(countries=["Uganda"]), "level": 4}
    assert ladder.combine([a, b])["level"] == 5
    assert ladder.combine([a, {**b, "countries": ["Kenya"]}])["level"] == 4


def test_replicated_in_africa_needs_two_e3_african_evaluations():
    a = {**_ev(countries=["Ghana"]), "level": 3}
    b = {**_ev(countries=["Rwanda"]), "level": 3}
    c = {**_ev(countries=["India"]), "level": 3}
    assert ladder.combine([a, b])["replicated_in_africa"] is True
    assert ladder.combine([a, c])["replicated_in_africa"] is False


# --- posture gates and geography ------------------------------------------------------
def test_posture_gates():
    assert ladder.best_posture(5, False, GATES) == "adopt"
    assert ladder.best_posture(4, False, GATES) == "adopt"
    assert ladder.best_posture(3, False, GATES) == "adapt"
    assert ladder.best_posture(3, True, GATES) == "adopt"
    assert ladder.best_posture(2, True, GATES) == "watch"
    assert ladder.best_posture(1, False, GATES) == "watch"


def test_cap_posture_lowers_never_raises():
    assert ladder.cap_posture("adopt", 2, False, GATES) == "watch"
    assert ladder.cap_posture("adopt", 3, False, GATES) == "adapt"
    assert ladder.cap_posture("watch", 5, False, GATES) == "watch"
    assert ladder.cap_posture("enter", 4, False, GATES) == "adopt"


@pytest.mark.parametrize("countries,region,expected", [
    (["Ghana"], "", True), (["India", "Kenya"], "", True), (["India"], "", False),
    ([], "Africa", True), ([], "Global", None), ("Côte d'Ivoire; Niger", "", True),
])
def test_is_african(countries, region, expected):
    assert ladder.is_african(countries, region) is expected


def test_africa_first_entry_bar():
    geo = {"africa_min": 2, "non_africa_min": 3}
    assert ladder.entry_allowed(2, True, geo) and not ladder.entry_allowed(1, True, geo)
    assert ladder.entry_allowed(2, None, geo)
    assert not ladder.entry_allowed(2, False, geo) and ladder.entry_allowed(3, False, geo)


# --- in the pipeline ------------------------------------------------------------------
EVCFG = {"search_first": ["3ie"], "max_evaluations": 2, "gates": GATES,
         "geography": {"africa_min": 2, "non_africa_min": 3}}


@pytest.fixture
def graded_profile(monkeypatch):
    monkeypatch.setattr(config, "SPEC", {**spec.DEFAULT_SPEC, "profile": "yes", "evidence": EVCFG,
                                         "windows": {"program": {"from": 2023, "to": 2026},
                                                     "evaluation": {"from": 2015, "to": 2026}}})
    monkeypatch.setattr(config, "DRY_RUN", True)


def test_dry_run_evidence_is_capped_without_grounding(graded_profile):
    org = {"name": "LEAP Africa", "region": "Africa"}
    appr = {"name": "Youth leadership program", "what": "w", "url": "https://example.org/p.pdf", "year": "2024"}
    rec = asyncio.run(pipeline._evidence_for({}, org, appr, EVCFG))
    assert rec["level"] <= 2 and rec["label"] in ("E1", "E2")
    assert rec["posture_allowed"] == "watch"
    assert rec["evaluations"] and all(e["method_grounded"] is None for e in rec["evaluations"])


def test_grounded_evidence_reaches_e4_and_the_verifier_runs(graded_profile, monkeypatch):
    monkeypatch.setattr(sources, "quote_exact", lambda url, q: True)
    org = {"name": "LEAP Africa", "region": "Africa"}
    # a subject whose hash makes the mock return an RCT reading
    for name in ("Program A", "Program B", "Program C", "Program D"):
        appr = {"name": name, "what": "w", "url": "https://example.org/p.pdf", "year": "2024"}
        rec = asyncio.run(pipeline._evidence_for({}, org, appr, EVCFG))
        own = [e for e in rec["evaluations"] if e.get("own")]
        assert own and all(e["level"] <= 3 for e in own), "the program's own document is never independent"
        if rec["level"] >= 3:
            best = max(rec["evaluations"], key=lambda e: e["level"])
            assert "verification" in best
            return
    pytest.fail("no mock reading reached E3 or above")


def test_out_of_window_evaluations_are_skipped(graded_profile, monkeypatch):
    async def old(*a, **k):
        return [{"title": "Old study", "year": "2009", "evaluator": "x", "type": "impact evaluation",
                 "url": "https://example.org/old.pdf"}]
    monkeypatch.setattr(pipeline.agents, "find_evidence", old)
    rec = asyncio.run(pipeline._evidence_for({}, {"name": "O", "region": "Africa"},
                                             {"name": "P", "what": "w", "url": "", "year": "2024"}, EVCFG))
    assert rec["evaluations"] == [] and rec["note"] == "no independent evaluation found"


# --- review overrides and the stage-two gates -----------------------------------------
def test_reviewer_override_is_logged_and_recomputes_posture(graded_profile, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "REVIEW_DIR", tmp_path)
    monkeypatch.setattr(config, "WORK_DIR", tmp_path)
    row = {"name": "P", "org": "O", "verification": {"status": "verified"},
           "evidence_record": {"level": 2, "label": "E2", "replicated_in_africa": False,
                               "posture_allowed": "watch", "best": {}, "flags": []}}
    io_xlsx.write_longlist([row])
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "longlist.xlsx")
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header[-4:] == ["evidence level", "method sentence", "evaluation", "evidence flags"]
    ws.cell(row=2, column=header.index("evidence level") + 1, value="E4")
    wb.save(tmp_path / "longlist.xlsx")
    kept = io_xlsx.read_kept_longlist()
    rec = kept[0]["evidence_record"]
    assert rec["override"] == {"from": "E2", "to": "E4", "by": "review"}
    assert rec["posture_allowed"] == "adopt"
    p = io_xlsx.write_evidence_overrides(kept)
    assert "O: P, E2 to E4" in p.read_text()


def test_horizon_longlist_columns_unchanged(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SPEC", None)
    monkeypatch.setattr(config, "REVIEW_DIR", tmp_path)
    monkeypatch.setattr(config, "WORK_DIR", tmp_path)
    io_xlsx.write_longlist([{"name": "P", "org": "O", "verification": {}}])
    from openpyxl import load_workbook
    header = [c.value for c in load_workbook(tmp_path / "longlist.xlsx").active[1]]
    assert header[-1] == "source"


def test_stage_two_gates_lower_theme_postures():
    rows = [{"name": "A", "evidence_record": {"level": 2}}, {"name": "B", "evidence_record": {"level": 3}}]
    themes = [{"name": "T1", "posture": "adopt", "members": ["A", "B"]},
              {"name": "T2", "posture": "adopt", "members": ["A"]},
              {"name": "T3", "posture": "watch", "members": ["B"]}]
    pipeline.apply_posture_gates(rows, themes, GATES)
    assert [r["posture"] for r in rows] == ["watch", "adapt"]
    assert [(t["posture"], t["evidence_level"]) for t in themes] == [
        ("adapt", "E3"), ("watch", "E2"), ("watch", "E3")]
    assert "allows adapt" in themes[0]["posture_note"]


def test_read_evidence_keeps_the_models_level(monkeypatch):
    """Regression: the level was lowercased before an uppercase compare, so every
    reading fell to E1."""
    async def fake(**kw):
        return {"model_level": "E4", "method": "rct"}
    monkeypatch.setattr(pipeline.agents, "structured_call", fake)
    r = asyncio.run(pipeline.agents.read_evidence({}, {"name": "P"}, {"url": "u"}, "text"))
    assert r["model_level"] == "E4" and r["method"] == "rct"


def test_peer_reviewed_counts_as_independent():
    g = ladder.grade(_ev(evaluator="Harambee", independent=False, peer_reviewed=True),
                     implementer="Harambee", method_grounded=True)
    assert g["level"] == 4


def test_golden_eval_scoring(monkeypatch, tmp_path):
    """The accuracy check's arithmetic, with the model and the network stubbed."""
    import json as _json
    from scan import agents as _agents, evaluate
    golden = {"targets": {"exact_match_min": 0.9, "false_e4_plus_max": 0, "e3_plus_ungrounded_max": 0},
              "evaluations": [
                  {"id": "rct", "program": "P", "implementer": "Gov", "title": "T", "evaluator": "Uni",
                   "url": "https://x.org/a", "expected_level": 4},
                  {"id": "tracer", "program": "P", "implementer": "Gov", "title": "T", "evaluator": "Gov",
                   "url": "https://x.org/b", "expected_level": 2}]}
    path = tmp_path / "golden.json"
    path.write_text(_json.dumps(golden))
    monkeypatch.setattr(sources, "fetch_text", lambda url, n=None: "text")
    monkeypatch.setattr(sources, "quote_exact", lambda url, q: True)
    readings = {
        "https://x.org/a": {"method": "rct", "model_level": "E4", "method_quote": "Participants were randomly assigned.",
                            "independent": True, "outcome_type": "outcomes", "program_matches": True},
        # the trap: data checks were randomized, the design is before and after
        "https://x.org/b": {"method": "before_after", "model_level": "E4",
                            "method_quote": "We ran randomized response verifications.",
                            "independent": False, "outcome_type": "outcomes", "program_matches": True},
    }

    async def fake_read(ctx, appr, ev, doc):
        return {"outcome_quote": "", "peer_reviewed": False, **readings[ev["url"]]}
    monkeypatch.setattr(_agents, "read_evidence", fake_read)
    res = asyncio.run(evaluate.evidence_eval(path))
    assert [r["graded"] for r in res["rows"]] == [4, 2]
    assert res["passed"] and res["false_e4_plus"] == []
    assert "PASS" in evaluate.evidence_report(res)


def test_golden_file_is_well_formed():
    import json as _json
    from pathlib import Path
    g = _json.loads((Path(__file__).resolve().parent.parent / "profiles" / "yes" / "golden.json").read_text())
    ev = g["evaluations"]
    assert 12 <= len(ev) <= 15
    assert {e["expected_level"] for e in ev} == {1, 2, 3, 4, 5}
    assert all(e["url"].startswith("https://") for e in ev)



def test_trap_randomized_data_checks_stay_low_when_the_model_reads_the_design():
    """'Randomized response verifications' match the E4 words, so the code alone would
    say E4. The model's method and level are ceilings too, so the lower level stands."""
    g = ladder.grade(_ev(method="before_after", model_level="E2",
                         method_quote="A QA team conducted randomized response verifications."),
                     implementer="GIZ", method_grounded=True)
    assert g["level"] == 2


def test_cap_note_only_when_it_lowers():
    g = ladder.grade(_ev(method="descriptive", model_level="E1", method_quote=""), implementer="X",
                     method_grounded=None)
    assert g["level"] == 1 and g["caps"] == []



def test_a_failed_evidence_search_is_an_error_not_e1(graded_profile, monkeypatch):
    async def boom(*a, **k):
        raise RuntimeError("Error code: 402, credits")
    monkeypatch.setattr(pipeline.agents, "find_evidence", boom)
    with pytest.raises(RuntimeError):
        asyncio.run(pipeline._evidence_for({}, {"name": "O", "region": "Africa"},
                                           {"name": "P", "what": "w", "url": "", "year": "2024"}, EVCFG))


def test_orgs_with_errors_are_scanned_again():
    assert not pipeline._scanned_cleanly(None)
    assert not pipeline._scanned_cleanly({"error": "boom", "dropped": []})
    assert not pipeline._scanned_cleanly({"dropped": [{"stage": "error", "reason": "402"}]})
    assert pipeline._scanned_cleanly({"dropped": [{"stage": "below the evidence bar"}]})



def test_empty_first_search_gets_a_broader_retry(graded_profile, monkeypatch):
    calls = []

    async def find(ctx, org, appr, sf, hint=""):
        calls.append(hint)
        if not hint:
            return []
        return [{"title": "Impact evaluation", "year": "2020", "evaluator": "Uni", "type": "impact evaluation",
                 "url": "https://example.org/eval.pdf"}]
    monkeypatch.setattr(pipeline.agents, "find_evidence", find)
    rec = asyncio.run(pipeline._evidence_for({}, {"name": "O", "region": "Africa"},
                                             {"name": "P", "what": "w", "url": "", "year": "2024"}, EVCFG))
    assert len(calls) == 2 and calls[1] and rec["evaluations"]


def test_below_the_bar_stays_in_the_longlist_unticked(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "REVIEW_DIR", tmp_path)
    monkeypatch.setattr(config, "WORK_DIR", tmp_path)
    monkeypatch.setattr(config, "SPEC", {**spec.DEFAULT_SPEC, "evidence": EVCFG})
    io_xlsx.write_longlist([{"name": "P", "org": "O", "verification": {}, "keep_default": "N",
                             "evidence_record": {"label": "E1", "flags": ["below the evidence bar: E1"]}}])
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "longlist.xlsx").active
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert row[header.index("keep")] == "N"
    assert "below the evidence bar" in row[header.index("evidence flags")]


def test_evaluation_on_the_implementers_own_site_is_self_published():
    g = ladder.grade(_ev(url="https://www.experienceeducate.org/blog/new-rct"), implementer="Educate!",
                     method_grounded=True, implementer_site="https://www.experienceeducate.org/")
    assert g["level"] == 3 and "published by the implementer" in g["caps"][0]
    peer = ladder.grade(_ev(url="https://www.harambee.co.za/wp-content/uploads/aer.pdf", peer_reviewed=True),
                        implementer="Harambee", method_grounded=True, implementer_site="https://www.harambee.co.za/")
    assert peer["level"] == 4, "a peer-reviewed paper hosted by the implementer still counts"


def test_e5_needs_trials_in_different_places():
    a = {**_ev(countries=["Uganda"]), "level": 4}
    b = {**_ev(countries=["Uganda", "Kenya"]), "level": 4}
    assert ladder.combine([a, b])["level"] == 4
    c = {**_ev(countries=["Rwanda"]), "level": 4}
    assert ladder.combine([a, c])["level"] == 5


def test_search_recall_matching(monkeypatch, tmp_path):
    import json as _json
    from scan import agents as _agents, evaluate
    golden = {"targets": {"search_recall_min": 0.5}, "evaluations": [
        {"id": "a", "program": "P", "implementer": "I", "expected_level": 4,
         "title": "Job Search and Hiring with Limited Information about Workseekers' Skills", "url": "https://x.org/aer.pdf"},
        {"id": "b", "program": "Q", "implementer": "J", "expected_level": 4,
         "title": "Generating Skilled Self-Employment in Developing Countries", "url": "https://y.org/qje.pdf"},
        {"id": "c", "program": "R", "implementer": "K", "expected_level": 1, "title": "t", "url": "https://z.org"}]}
    path = tmp_path / "g.json"
    path.write_text(_json.dumps(golden))

    async def find(ctx, org, appr, sf, hint=""):
        if appr["name"] == "P":
            return [{"title": "Job search and hiring with limited information about workseekers skills (AER)",
                     "url": "https://aeaweb.org/x"}]
        return [{"title": "Unrelated brief", "url": "https://w.org"}]
    monkeypatch.setattr(_agents, "find_evidence", find)
    res = asyncio.run(evaluate.search_recall(path))
    assert [(r["id"], r["found"]) for r in res["rows"]] == [("a", True), ("b", False)]
    assert res["passed"] and "1 of 2" in evaluate.recall_report(res)
