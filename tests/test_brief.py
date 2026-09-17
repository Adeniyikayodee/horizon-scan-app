"""B5: the plain-language brief checks, the rewrite pass, and the options workbook."""
from __future__ import annotations

import asyncio

import pytest
from openpyxl import load_workbook

from scan import agents, config, io_xlsx, plain, spec

CHECKS = {"max_grade": 9, "max_sentence_words": 20, "target_words": 0}


# --- antithesis -----------------------------------------------------------------------
@pytest.mark.parametrize("sentence", [
    "It is not just training, but real jobs.",
    "The program not only trains young people but also places them.",
    "This is not about skills, it is about jobs.",
    "Less talk, more action for young people.",
    "It is more than just a grant.",
    "Schools should teach trades rather than theory.",
    "Pay for results instead of activities.",
    "The fund is not a loan, but a grant.",
])
def test_antithesis_is_caught(sentence):
    assert plain.antithesis_hits(sentence), sentence


@pytest.mark.parametrize("sentence", [
    "The program trains young people and helps them find work.",
    "Employment rose by 12 percent in two years.",
    "Most programs did not measure earnings.",
])
def test_plain_statements_pass(sentence):
    assert plain.antithesis_hits(sentence) == []


# --- reading level and style ----------------------------------------------------------
def test_reading_grade_separates_plain_from_dense():
    easy = "The program helps young people find jobs, pays for their training, and works in Ghana."
    hard = ("Institutionalized multistakeholder operationalization of intergenerational employability "
            "interventions necessitates comprehensive organizational capacity strengthening.")
    assert plain.reading_grade(easy) <= 9 < plain.reading_grade(hard)


def test_headings_tables_and_links_are_not_prose():
    md = "# A Very Long Institutional Heading\n| Col | E4 |\n|---|---|\nRead [the study](https://x.org/a-b)."
    assert plain._prose(md) == "Read the study."


def test_acronyms_need_spelling_out_once():
    assert plain.acronym_hits("The African Center for Economic Transformation (ACET) works. ACET helps.") == []
    assert plain.acronym_hits("ACET helps. TVET matters.") == [
        "acronym not spelled out at first use: ACET", "acronym not spelled out at first use: TVET"]
    assert plain.acronym_hits("The program reached E4 on the scale.") == []


def test_house_style_numbers_percent_and_dashes():
    hits = plain.house_hits("We found 5 designs. Rates rose 12% — fast.")
    assert any("zero to nine" in h for h in hits)
    assert 'use "percent" in prose, not %' in hits
    assert "dash: use commas or rewrite" in hits
    assert plain.house_hits("We found five designs on June 5, 2026, rated E4, in 12 countries.") == []


def test_word_target_tolerance():
    body = "# T\n\n" + "Young people find work. " * 800
    assert any("the target is 2,700" in i for i in plain.check(body, {**CHECKS, "target_words": 2700}))
    ok = "# T\n\n" + "Young people find work. " * 675
    assert not any("target" in i for i in plain.check(ok, {**CHECKS, "target_words": 2700}))


def test_memo_ceiling():
    sp = {"memo": {"min_words": 10, "max_words": 20, "sections": []}}
    assert "ceiling" in spec.memo_shortfall("word " * 25, sp)
    assert spec.memo_shortfall("word " * 15, sp) == ""


# --- the rewrite pass -----------------------------------------------------------------
@pytest.fixture
def brief_profile(monkeypatch):
    monkeypatch.setattr(config, "SPEC", {**spec.DEFAULT_SPEC, "brief_checks": CHECKS,
                                         "memo": {"min_words": 1, "sections": [{"heading": "What we found"}]}})


def test_rewrite_kept_when_it_fixes_issues(brief_profile, monkeypatch):
    async def draft(ctx, themes, extra=""):
        return {"memo_markdown": "# T\n\n## What we found\n\nIt is not just training, but jobs.",
                "scorecard_intro": "i"}

    async def rewrite(ctx, md, issues):
        return "# T\n\n## What we found\n\nThe program trains young people for jobs."
    monkeypatch.setattr(agents, "_synthesize_draft", draft)
    monkeypatch.setattr(agents, "rewrite_brief", rewrite)
    out = asyncio.run(agents.synthesize({}, []))
    assert "trains young people" in out["memo_markdown"] and out["plain_issues"] == []


def test_rewrite_rejected_when_it_drops_a_section(brief_profile, monkeypatch):
    async def draft(ctx, themes, extra=""):
        return {"memo_markdown": "# T\n\n## What we found\n\nIt is not just training, but jobs.",
                "scorecard_intro": "i"}

    async def rewrite(ctx, md, issues):
        return "# T\n\nThe program trains young people for jobs."
    monkeypatch.setattr(agents, "_synthesize_draft", draft)
    monkeypatch.setattr(agents, "rewrite_brief", rewrite)
    out = asyncio.run(agents.synthesize({}, []))
    assert "not just training" in out["memo_markdown"]
    assert any("not just X but Y" in i for i in out["plain_issues"])


def test_no_checks_means_no_rewrite(monkeypatch):
    monkeypatch.setattr(config, "SPEC", None)
    called = []

    async def draft(ctx, themes, extra=""):
        return {"memo_markdown": "It is not just training, but jobs.", "scorecard_intro": ""}

    async def rewrite(*a):
        called.append(1)
        return ""
    monkeypatch.setattr(agents, "_synthesize_draft", draft)
    monkeypatch.setattr(agents, "rewrite_brief", rewrite)
    out = asyncio.run(agents.synthesize({}, []))
    assert called == [] and "plain_issues" not in out


# --- the options workbook -------------------------------------------------------------
def test_options_and_funder_map(tmp_path):
    rows = [
        {"name": "Opt A", "org": "Org", "posture": "adapt", "what": "w", "design_features": "Cash plus coaching",
         "target_group": "Young women", "countries": "Ghana", "delivery_partner": "Partner", "cost_data": "",
         "funders": "Fund X", "url": "https://org.org/p", "verification": {"status": "verified"},
         "score": {"reason_inclusion": "Reaches young women.", "reason_transferability": "Runs in Ghana.",
                   "reason_acet_role": "Ministry uptake."},
         "evidence_record": {"level": 3, "label": "E3", "funders": ["Fund Y"],
                             "best": {"title": "Eval", "url": "https://eval.org/e", "year": "2021",
                                      "effect_summary": "Employment rose — a lot.", "cost_per_outcome": "not found"}}},
        {"name": "Opt B", "org": "Org", "posture": "watch", "evidence_record": {"level": 1, "label": "E1",
                                                                              "note": "no independent evaluation found"}},
    ]
    themes = [{"name": "Entrepreneurship and youth-led business", "tag": "current", "members": ["Opt A", "Opt B"]}]
    fmap = [{"name": "Fund X", "type": "Foundation", "source": "roster",
             "fit": {"score": 9, "themes_matched": ["Entrepreneurship and youth-led business"], "countries_matched": ["ghana"]},
             "strategy": {"value": "Young Africa Works", "period": "2018-2030", "grounded": True, "url": "https://f.org/s"},
             "instruments": {"value": ["grants"]}, "size": {"value": "not found"}, "eligibility": {"value": "unclear"},
             "calls": [{"title": "Call", "deadline": "December 1, 2026", "url": "https://f.org/c"}], "dropped_fields": []},
            {"name": "Fund Z", "type": "", "source": "named on 1 option(s)", "error": "timeout"}]
    path, hits = io_xlsx.write_options(rows, themes, fmap, tmp_path / "options.xlsx")
    wb = load_workbook(path)
    assert wb.properties.creator == "Kayode Adeniyi"
    opts = list(wb["Options"].iter_rows(values_only=True))
    assert list(opts[0]) == io_xlsx.OPTION_COLUMNS
    a = dict(zip(opts[0], opts[1]))
    assert a["Option"] == "Opt A" and a["Evidence level"] == "E3" and a["Posture"] == "adapt"
    assert a["Relation"] == "current" and a["Funders"] == "Fund Y; Fund X"
    assert a["Outcome and effect"] == "Employment rose, a lot." and a["Cost per outcome"] == "not found"
    assert a["Transfer note"] == "Runs in Ghana." and a["ACET's role"] == "Ministry uptake."
    b = dict(zip(opts[0], opts[2]))
    assert b["Evaluation"] == "no independent evaluation found"
    fm = list(wb["Funder map"].iter_rows(values_only=True))
    assert list(fm[0]) == io_xlsx.FUNDER_COLUMNS
    x = dict(zip(fm[0], fm[1]))
    assert x["Fit score"] == 9 and x["Open calls"] == "Call, December 1, 2026, https://f.org/c"
    assert dict(zip(fm[0], fm[2]))["Not confirmed"] == "could not be read"
    assert hits == []



def test_choppy_short_sentences_are_caught():
    text = "The program works. It trains young people. It runs in Ghana, Kenya, and Rwanda for young women."
    hits = plain.short_sentence_hits(text, 7)
    assert len(hits) == 2 and "join it to the next" in hits[0]
    joined = ("The program works and trains young people; it runs in Ghana, Kenya, and Rwanda, "
              "where it reaches young women in rural areas.")
    assert plain.short_sentence_hits(joined, 7) == []


def test_min_sentence_words_is_part_of_the_check():
    md = "# T\n\n## What we found\n\nThe program works. It trains young people."
    assert any("short sentence" in i for i in plain.check(md, {"min_sentence_words": 7}))
    assert not any("short sentence" in i for i in plain.check(md, {}))


def test_names_do_not_inflate_the_reading_grade():
    plain_text = "The German agency works with young people in Rwanda and Uganda on training for jobs."
    named = ("The Deutsche Gesellschaft für Internationale Zusammenarbeit works with young people in Rwanda "
             "and Uganda on training for jobs.")
    assert plain.reading_grade(named) - plain.reading_grade(plain_text) < 2


def test_method_talk_is_caught():
    assert plain.method_talk_hits("This reflects a gap in our current search methodology.")
    assert plain.method_talk_hits("None of these funders had open calls identified during this scan period.")
    assert plain.method_talk_hits("Young women earned more after the training.") == []


def test_every_theme_must_be_named():
    md = "## Other options\n\nTVET at secondary level had one option. Secondary education had two."
    hits = plain.theme_coverage_hits(md, ["TVET at secondary level", "Secondary education", "Financing education"])
    assert hits == ["theme not covered by name: Financing education"]


def test_an_organizations_own_acronym_name_is_allowed():
    assert plain.acronym_hits("GIF funds innovation.", {"GIF"}) == []
    assert plain.acronym_hits("GIF funds innovation.") == ["acronym not spelled out at first use: GIF"]
