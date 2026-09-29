"""B6: the YES profile loads, runs end to end in test mode, and keeps its rules."""
from __future__ import annotations

import asyncio
import json
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from scan import config, io_xlsx, pipeline, spec

PROFILE = Path(__file__).resolve().parent.parent / "profiles" / "yes"


def test_profile_is_complete():
    sp = json.loads((PROFILE / "profile.json").read_text())
    seed = sp["themes_seed"]
    assert len(seed) >= 10 and len({t["name"] for t in seed}) == len(seed)
    assert sum(1 for t in seed if t["relation"] == "current") == 6
    assert all(t["relation"] in sp["tags"] for t in seed)
    keys = [c["key"] for c in sp["criteria"]]
    assert keys[:2] == ["evidence_strength", "research_gap"] and "acet_role" in keys
    # the criteria carry equal weight, so no dimension is searched less deeply
    assert len({c["weight"] for c in sp["criteria"]}) == 1
    assert sp["evidence"]["score_key"] == "evidence_strength"
    # the brief's spine is where the research runs out
    assert sp["memo"]["sections"][2]["heading"] == "The research gaps to take on"
    assert sp["memo"]["min_words"] < sp["brief_checks"]["target_words"] < sp["memo"]["max_words"]
    for f in ("themes.md", "output_spec.md", "policy.md", "exemplar.md"):
        assert (PROFILE / "context" / f).exists()
    orgs = io_xlsx.read_orgs(PROFILE / "organizations.xlsx")
    assert len(orgs) >= 40 and all(o.get("website", "").startswith("https://") for o in orgs)
    # a gap scan has to read the bodies that produce evidence, not only those that spend it
    kinds = " ".join(o.get("type", "").lower() for o in orgs)
    assert "research" in kinds and "evidence platform" in kinds


def test_evidence_mark_is_set_from_the_level():
    appr = {"score": {"evidence_strength": "strong", "reason_evidence_strength": "model"}}
    pipeline._set_evidence_mark(appr, {"level": 2, "label": "E2", "best": {"title": "Tracer study"}},
                                {"score_key": "evidence_strength"})
    assert appr["score"]["evidence_strength"] == "weak"
    assert appr["score"]["reason_evidence_strength"] == "Set from the evidence level, E2: Tracer study."


def test_press_cannot_verify(monkeypatch):
    monkeypatch.setattr(config, "SPEC", {"unverifiable_source_types": ["press"]})
    row = {"source_type": "press", "verification": {"status": "verified", "note": ""}}
    pipeline._hold_unverifiable_source(row)
    assert row["verification"]["status"] == "partial"


def test_yes_dry_run_end_to_end(monkeypatch):
    base = Path(tempfile.mkdtemp(prefix="yes-dry-"))
    pdir = base / "profiles" / "yes"
    shutil.copytree(PROFILE, pdir, ignore=shutil.ignore_patterns("run"))
    monkeypatch.setattr(config, "PROFILES_DIR", base / "profiles")
    monkeypatch.setattr(config, "ROOT", base)
    for k in ("SPEC", "WORK_DIR", "ORGS_WORK", "REVIEW_DIR", "OUT_DIR", "MANIFEST", "ORG_SHEET", "DRY_RUN"):
        monkeypatch.setattr(config, k, getattr(config, k))
    config.DRY_RUN = True
    config.use_profile("yes")
    asyncio.run(pipeline.run_stage1())
    asyncio.run(pipeline.run_stage2())
    out = config.OUT_DIR
    assert (out / "yes_scan_brief.md").exists() and (out / "yes_scan_brief.docx").exists()
    wb = load_workbook(out / f"{config.active_spec()['deliverables']['options']}.xlsx")
    assert wb.sheetnames == ["Options", "Research gaps", "Coverage map", "Funder map"]
    cov = list(wb["Coverage map"].iter_rows(values_only=True))
    assert cov[0][:4] == ("Theme", "Posture", "Options", "Best evidence") and "Ghana" in cov[0]
    gaps = list(wb["Research gaps"].iter_rows(values_only=True))
    assert gaps[0][0] == "Research gap" and len(gaps) > 1
    rows = list(wb["Options"].iter_rows(values_only=True))
    header = rows[0]
    themes = {r[header.index("Theme")] for r in rows[1:]}
    allowed = {t["name"] for t in config.active_spec()["themes_seed"]} | {None, ""}
    assert themes <= allowed, "every option sits in a theme from the fixed list, or awaits review"
    postures = {r[header.index("Posture")] for r in rows[1:]}
    assert postures <= {"use", "scope", "commission"}
    # nothing in test mode can be grounded, so every question reads as uncovered
    assert postures == {"commission"}
    assert len(list(wb["Funder map"].iter_rows())) - 1 >= 21
    assert not (out / "theme_scorecard.xlsx").exists()


def test_one_report_per_candidate_in_the_yes_profile():
    reports = [{"title": "Employment-focused subject in secondary schools evidence", "url": "https://e.org/evidence"}]
    a = {"name": "Employment-focused subject", "one_liner": "secondary schools"}
    b = {"name": "Livelihood bootcamps", "one_liner": "employment-focused support for young women outside secondary schools"}
    used: set = set()
    first = pipeline._best_report(a, reports, used)
    used.add(first["url"])
    assert pipeline._best_report(b, reports, used) is None
    assert pipeline._best_report(b, reports, None) is not None, "the horizon scan keeps today's matching"
